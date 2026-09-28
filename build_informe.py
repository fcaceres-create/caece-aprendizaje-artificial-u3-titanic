"""Genera Informe_Titanic.docx a partir de los resultados en output/ y las figuras en figs/.
Ejecutar después del notebook:  python build_informe.py
"""
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

comp = pd.read_csv("output/comparacion_modelos.csv", index_col=0)
tun = pd.read_csv("output/tuning.csv", index_col=0)
ho = pd.read_csv("output/holdout.csv", index_col=0)
imp = pd.read_csv("output/importancia_permutacion.csv", index_col=0)
sub = pd.read_csv("output/submission.csv")
mejor = tun["Acc CV"].idxmax()

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)

def h(t, n=1): doc.add_heading(t, level=n)
def p(t, bold_prefix=None):
    par = doc.add_paragraph()
    if bold_prefix: par.add_run(bold_prefix).bold = True
    par.add_run(t); par.paragraph_format.space_after = Pt(6)
    return par
def bullets(items):
    for it in items:
        par = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            par.add_run(it[0]).bold = True; par.add_run(it[1])
        else:
            par.add_run(it)
def fig(nombre, ancho=15.5, caption=None):
    doc.add_picture(f"figs/{nombre}.png", width=Cm(ancho))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        c = doc.add_paragraph(caption); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].italic = True; c.runs[0].font.size = Pt(9); c.runs[0].font.color.rgb = RGBColor(90, 90, 90)
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)
def tabla(df, index=True, fmt="{:.3f}"):
    cols = ([df.index.name or ""] if index else []) + list(df.columns)
    t = doc.add_table(rows=1, cols=len(cols)); t.style = "Light Grid Accent 1"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = str(c)
    for idx, row in df.iterrows():
        cells = t.add_row().cells; vals = ([idx] if index else []) + list(row)
        for i, v in enumerate(vals):
            cells[i].text = fmt.format(v) if isinstance(v, float) and pd.notna(v) else ("—" if isinstance(v, float) else str(v))
    for r in t.rows:
        for c in r.cells:
            for par in c.paragraphs:
                for run in par.runs: run.font.size = Pt(9.5)
    doc.add_paragraph()
    return t

# ---------------------------------------------------------------- Portada
doc.add_heading("Predicción de supervivencia en el Titanic", 0)
p("Trabajo práctico — Materia: ______________________")
p("Integrantes: ______________________________________________")
p("Herramienta: Python 3 (pandas, scikit-learn, matplotlib/seaborn) en un Jupyter Notebook.")

h("1. Objetivo y planteo del problema")
p("El objetivo es predecir si un pasajero del Titanic sobrevivió (Survived = 1) o no (Survived = 0) "
  "a partir de sus características, e identificar cuáles de esas características determinan una mayor "
  "probabilidad de sobrevivir. Es un problema de clasificación binaria supervisada.")
p("Se usan los datos de la competencia de Kaggle: train.csv (891 pasajeros, con la etiqueta) y test.csv "
  "(418 pasajeros, sin etiqueta, sobre el que se generan las predicciones para enviar). La métrica "
  "oficial es la accuracy (proporción de aciertos); además se reportan ROC-AUC y F1 porque las clases "
  "no están balanceadas (38% sobrevivió).")

h("2. Metodología general")
bullets([
    ("Exploración (EDA): ", "entender las variables, los faltantes y la relación de cada una con la supervivencia."),
    ("Preparación: ", "imputar faltantes y crear variables nuevas, todo dentro de un Pipeline para evitar fuga de información."),
    ("Modelado: ", "comparar seis algoritmos contra dos líneas base con validación cruzada estratificada de 10 folds."),
    ("Ajuste: ", "búsqueda de hiperparámetros (GridSearchCV) para los tres mejores candidatos de naturaleza distinta."),
    ("Evaluación: ", "medición final sobre un 20% de hold-out que no participó en ninguna decisión."),
    ("Interpretación: ", "coeficientes de regresión logística, importancia por permutación y un árbol de reglas."),
    ("Entrega: ", "reentrenamiento con los 891 registros y generación de submission.csv."),
])
p("Se fijó una semilla aleatoria (SEED = 42) en todas las particiones y modelos para que los resultados "
  "sean exactamente reproducibles.")

h("3. Análisis exploratorio")
h("3.1 Calidad de los datos", 2)
tabla(pd.DataFrame({"Faltantes train": [177, 687, 2, 0], "% train": ["19,9%", "77,1%", "0,2%", "0%"],
                    "Faltantes test": [86, 327, 0, 1]},
                   index=pd.Index(["Age", "Cabin", "Embarked", "Fare"], name="Variable")))
p("Age falta en uno de cada cinco pasajeros y es una variable relevante, por lo que hay que imputarla con "
  "cuidado. Cabin falta en el 77%: imputar el valor exacto sería inventar datos, así que sólo se aprovecha "
  "si hay cabina registrada y la cubierta (letra). Embarked y Fare tienen faltantes mínimos.")

h("3.2 Hallazgos principales", 2)
fig("01_sexo_clase_puerto", caption="Figura 1. Tasa de supervivencia por sexo, clase y puerto (IC 95%).")
bullets([
    ("Sexo: ", "es el factor dominante. Sobrevivió el 74% de las mujeres y el 19% de los hombres."),
    ("Clase: ", "63% en 1ra, 47% en 2da y 24% en 3ra. Interactúa con el sexo: las mujeres de 1ra y 2da "
     "sobrevivieron en un 97% y 92%, las de 3ra sólo en un 50%."),
    ("Puerto: ", "Cherbourg muestra más supervivencia, pero allí el 51% embarcó en 1ra clase. Es un efecto "
     "confundido con la clase, no un efecto propio del puerto."),
])
fig("02_edad_tarifa", caption="Figura 2. Distribución de edad y tarifa según supervivencia.")
p("Los niños pequeños sobreviven más (\"mujeres y niños primero\"). La tarifa tiene una distribución muy "
  "asimétrica (máximo 512) y refleja la clase; por eso se transforma con logaritmo.")
fig("03_familia", ancho=11, caption="Figura 3. Supervivencia según tamaño de familia a bordo.")
p("La relación con el tamaño de familia no es lineal: quienes viajaban solos o en familias de 5 o más "
  "sobreviven menos que las familias de 2 a 4. Un modelo lineal no captura esa forma si se usa el número "
  "tal cual, lo que justifica agrupar la variable.")
p("El título que figura en el nombre (Mr, Mrs, Miss, Master, etc.) resultó muy informativo: \"Master\" "
  "(niños varones) sobrevive un 57% frente a un 16% de \"Mr\". Así se distingue a los niños varones de los "
  "hombres adultos, algo que la variable Sex sola no permite.")

h("4. Preparación de los datos y creación de variables")
p("todo el preprocesamiento se implementó como un transformador propio (TitanicFeatures) dentro de un "
  "Pipeline de scikit-learn. Los valores que se \"aprenden\" de los datos (mediana de edad por título, "
  "mediana de tarifa por clase, moda del puerto, frecuencia de cada ticket) se calculan sólo con la parte "
  "de entrenamiento de cada fold. Si se imputara antes de dividir, la información del conjunto de "
  "validación se filtraría al modelo (data leakage) y las métricas serían optimistas.",
  "Justificación del Pipeline: ")
tabla(pd.DataFrame([
    ("Title", "Texto entre la coma y el punto de Name; Mlle/Ms→Miss, Mme→Mrs, el resto de los poco frecuentes→Rare", "Resume sexo, edad y estatus social"),
    ("AgeImp", "Age imputada con la mediana del título", "Más precisa que la mediana global: un Master no tiene 28 años"),
    ("AgeMissing", "1 si Age era nula", "El faltante puede ser informativo en sí mismo"),
    ("IsChild", "AgeImp < 13", "Refleja la regla de niños primero"),
    ("FamilySize / FamilyGroup", "SibSp + Parch + 1; agrupado en Solo / Chica (2–4) / Grande (≥5)", "Captura la relación no lineal"),
    ("TicketGroup", "Pasajeros que comparten ticket (tope 5)", "Grupos que viajaban juntos sin ser familia"),
    ("HasCabin / Deck", "Si hay cabina registrada; primera letra (U = desconocida)", "Aprovechar Cabin sin inventar datos"),
    ("LogFare", "log(1 + Fare); el faltante se completa con la mediana de su clase", "Reducir la asimetría y el efecto de outliers"),
], columns=["Variable", "Construcción", "Justificación"]).set_index("Variable"))
p("Luego las variables numéricas se estandarizan (media 0, desvío 1), necesario para regresión logística, "
  "SVM y KNN, que son sensibles a la escala, y las categóricas (Sex, Pclass, Title, Embarked, FamilyGroup, "
  "Deck) se codifican con one-hot. Pclass se trata como categórica porque la distancia entre 1ra y 2da no "
  "tiene por qué ser la misma que entre 2da y 3ra. Se descartan PassengerId (identificador sin información) "
  "y las columnas de texto crudo Name, Ticket y Cabin, ya resumidas en las variables anteriores.")

h("5. Modelado")
h("5.1 Esquema de validación", 2)
p("Se separó un 20% estratificado (179 filas) como hold-out, reservado para la evaluación final. Sobre el "
  "80% restante (712 filas) se comparan los modelos con validación cruzada estratificada de 10 folds. Con "
  "tan pocos datos, una sola partición train/validación daría una estimación muy ruidosa; la validación "
  "cruzada promedia 10 estimaciones y además muestra su variabilidad. La estratificación mantiene en cada "
  "fold la misma proporción de sobrevivientes.")
h("5.2 Líneas base", 2)
p("Antes de entrenar modelos se fijan dos referencias: predecir siempre \"no sobrevive\" logra un 61,7% de "
  "accuracy, y la regla \"sobreviven todas las mujeres\" un 78,9%. Un modelo sólo aporta valor si supera "
  "claramente esta segunda regla.")
h("5.3 Comparación de algoritmos", 2)
p("Se eligieron algoritmos de familias diferentes: lineal (regresión logística), basado en distancias "
  "(KNN), de márgenes (SVM con kernel RBF), árbol de decisión y dos ensambles (Random Forest por bagging y "
  "Gradient Boosting por boosting).")
tabla(comp[["Acc train", "Acc CV", "± desv", "AUC CV", "F1 CV", "Brecha (sobreajuste)"]])
fig("05_comparacion_modelos", caption="Figura 4. Accuracy de validación (barras) frente a entrenamiento (puntos).")
p("El árbol de decisión sin restricciones y Random Forest alcanzan casi 99% en entrenamiento pero "
  "bastante menos en validación: memorizan los datos. El árbol individual ni siquiera supera la regla "
  "\"mujeres viven\". La columna Brecha cuantifica el sobreajuste y motivó el ajuste de hiperparámetros.",
  "Problema detectado, sobreajuste: ")

h("5.4 Ajuste de hiperparámetros", 2)
p("Con GridSearchCV (misma validación cruzada) se ajustaron un modelo lineal y los dos ensambles. En los "
  "árboles se exploraron parámetros que limitan la complejidad (profundidad máxima, mínimo de muestras por "
  "hoja, fracción de variables); en la regresión logística la regularización C. SVM, con resultados "
  "similares a la regresión logística, no se ajustó para no multiplicar pruebas sobre un dataset tan chico.")
tun_show = tun.copy(); tun_show["Mejores parámetros"] = tun_show["Mejores parámetros"].str.replace("'", "")
tabla(tun_show)
p(f"El ajuste redujo drásticamente el sobreajuste: Random Forest pasó de una brecha de "
  f"{comp.loc['Random Forest', 'Brecha (sobreajuste)']:.2f} a "
  f"{tun.loc['Random Forest', 'Acc train'] - tun.loc['Random Forest', 'Acc CV']:.2f} y su accuracy de "
  f"validación subió de {comp.loc['Random Forest', 'Acc CV']:.3f} a {tun.loc['Random Forest', 'Acc CV']:.3f}. "
  f"Los mejores valores eligen árboles poco profundos: con estos datos, modelos más simples generalizan mejor.")

h("6. Evaluación final (hold-out)")
p(f"El modelo final se eligió por su accuracy en validación cruzada ({mejor}), no por el hold-out, para no "
  "usar el conjunto de prueba en la toma de decisiones. Luego se midieron todos sobre el hold-out:")
tabla(ho)
p("Los tres modelos superan a la línea base por 2 a 6 puntos. Las diferencias entre ellos (1 a 3 puntos, "
  "o sea 2 a 6 pasajeros de 179) están dentro del margen de ruido esperable con tan pocas observaciones; "
  "es honesto decir que regresión logística y Random Forest rinden prácticamente igual.")
fig("07_confusion", ancho=9, caption=f"Figura 5. Matriz de confusión de {mejor} en el hold-out.")
p("El error más frecuente son los sobrevivientes predichos como fallecidos (recall de la clase positiva "
  "≈ 74%). Suelen ser hombres adultos que sobrevivieron o mujeres de 3ra clase que no: casos que "
  "contradicen el patrón general y que con las variables disponibles no se pueden separar.")

h("7. ¿Qué características determinan la supervivencia?")
fig("08_coeficientes_lr", ancho=12, caption="Figura 6. Coeficientes de la regresión logística (variables estandarizadas).")
fig("09_importancia_permutacion", ancho=13, caption=f"Figura 7. Importancia por permutación de {mejor} en el hold-out.")
p("Los dos métodos coinciden en el orden de importancia:")
bullets([
    ("Sexo y título: ", "ser mujer casi duplica las chances (odds ratio ≈ 1,9); el título \"Mr\" las reduce a menos "
     "de la mitad. Al permutar Sex la accuracy cae unos 10 puntos, y al permutar Name (de donde sale Title) unos 6."),
    ("Clase socioeconómica: ", "3ra clase reduce la probabilidad; la tarifa y tener cabina registrada, que son "
     "reflejos del estatus, la aumentan."),
    ("Edad: ", "a mayor edad, menor probabilidad; ser niño (IsChild, Master) la aumenta."),
    ("Grupo familiar: ", "viajar en familias grandes reduce la supervivencia; las familias de 2–4 tienen ventaja."),
    ("Poco relevantes: ", "puerto de embarque, número de ticket y la cabina específica aportan prácticamente nada en "
     "datos nuevos (importancia ≈ 0)."),
])
fig("10_arbol_reglas", ancho=16.5, caption="Figura 8. Árbol de decisión de profundidad 3 (reglas interpretables).")
p("El árbol de profundidad 3 resume las reglas: quienes no son \"Mr\" y no viajan en 3ra clase sobreviven en un "
  "95%; los \"Mr\" sin cabina registrada y con tarifa baja mueren en un 90%. Este árbol tan simple obtiene "
  "un 83% en el hold-out, lo que confirma que pocas variables explican casi todo el fenómeno.")

h("8. Modelo final y predicción")
p(f"El modelo elegido ({mejor}) se reentrenó con los 891 pasajeros y predijo los 418 del test. El archivo "
  f"output/submission.csv tiene el formato que exige Kaggle (PassengerId, Survived); predice que sobrevive el "
  f"{sub.Survived.mean():.0%} de los pasajeros del test, en línea con el 38% del entrenamiento. También se "
  "generaron envíos para regresión logística y Gradient Boosting para compararlos en el leaderboard.")
p("Es esperable que el puntaje público de Kaggle quede algo por debajo del de validación (típicamente 0,77–0,80 "
  "para este enfoque), porque el leaderboard se calcula sobre una fracción del test, que es pequeño.")

h("9. Problemas encontrados y cómo se resolvieron")
bullets([
    ("Datos faltantes: ", "Age se imputó por título; Cabin se transformó en indicador + cubierta; Embarked con la moda, "
     "Fare con la mediana de su clase."),
    ("Variables de texto: ", "Name, Ticket y Cabin no sirven crudas; se extrajo de ellas información útil (Title, TicketGroup, Deck)."),
    ("Fuga de información: ", "se evitó metiendo toda la preparación en un Pipeline que se ajusta dentro de cada fold."),
    ("Sobreajuste: ", "detectado comparando la accuracy de entrenamiento con la de validación; se corrigió limitando la "
     "complejidad con búsqueda de hiperparámetros."),
    ("Dataset chico: ", "validación cruzada de 10 folds con desvío estándar; las diferencias pequeñas entre modelos no se "
     "tratan como significativas."),
    ("Variables confundidas: ", "el efecto del puerto de embarque se explica por la clase; no se interpretó como causal."),
    ("Desbalance moderado: ", "estratificación en todas las particiones y métricas complementarias (AUC, F1)."),
])

h("10. Conclusiones")
p("La supervivencia en el Titanic estuvo determinada principalmente por el sexo, la edad (niños) y la clase "
  "socioeconómica, lo que es coherente con el protocolo histórico de \"mujeres y niños primero\" y con el acceso "
  "más fácil a los botes desde las cubiertas superiores. Con una buena ingeniería de variables, modelos "
  "relativamente simples alcanzan alrededor de 83% de accuracy, unos 4–5 puntos por encima de la regla trivial "
  "basada sólo en el sexo. Modelos más complejos no mejoraron de forma significativa: con 891 registros el "
  "límite lo ponen los datos, no el algoritmo.")
p("Posibles mejoras: variables de supervivencia por grupo familiar o de ticket (\"si murió el resto de la familia\"), "
  "ensambles de varios modelos (voting/stacking) y validación cruzada repetida para comparar modelos con más precisión.")

h("11. Reproducibilidad")
bullets([
    "Instalar dependencias: pip install -r requirements.txt",
    "Colocar train.csv y test.csv de Kaggle en la carpeta data/ (ya incluidos).",
    "Ejecutar el notebook titanic.ipynb de principio a fin (o: jupyter nbconvert --to notebook --execute titanic.ipynb).",
    "Resultados: tablas y submission.csv en output/, gráficos en figs/. La semilla fija garantiza los mismos números.",
    "Opcional: python build_informe.py regenera este documento a partir de los resultados.",
])

doc.save("Informe_Titanic.docx")
print("Informe_Titanic.docx generado")
